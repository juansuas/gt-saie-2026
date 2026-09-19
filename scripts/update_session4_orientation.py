from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "dist/documentos/orientaciones/orientaciones-sesion-4.docx"


def add_bullet(doc: Document, lead: str, body: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(f"{lead}: ")
    run.bold = True
    paragraph.add_run(body)


doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Pt(54)
section.left_margin = section.right_margin = Pt(62)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(20)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
title_border = styles["Title"].element.pPr.find(qn("w:pBdr"))
if title_border is not None:
    styles["Title"].element.pPr.remove(title_border)
for style_name, size in (("Heading 1", 14), ("Heading 2", 11.5)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12 if style_name == "Heading 1" else 10)
    style.paragraph_format.space_after = Pt(5 if style_name == "Heading 1" else 4)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("GT 1 Usos de información estadística y acceso a datos públicos")
subtitle = doc.add_paragraph()
subtitle.add_run("Sesión 4 Ejes para la discusión").bold = True
doc.add_paragraph("Jueves de 14:30 a 16:30")
doc.add_paragraph("Moderación a cargo de Delia González")

doc.add_heading("Trabajos presentados", level=1)
groups = [
    ("Grupo 1", "Cecilia Adrogué y Eugenia Orlicki"),
    ("Grupo 2", "Ivana Templado, Lara Encinas y Stefano Pistoia"),
    ("Grupo 3", "Mirta Torres y Diego Born"),
    ("Grupo 4", "Jimena Macció, Sonia Susini y Flavio Guberman"),
]
for group, authors in groups:
    add_bullet(doc, group, authors)

doc.add_heading("Dinámica propuesta para la presentación de trabajos", level=1)
add_bullet(doc, "Primer bloque", "Se presentan los trabajos de los grupos 1 y 2. Cada grupo dispondrá de aproximadamente 15 minutos para la presentación y se destinarán 30 minutos para el intercambio.")
add_bullet(doc, "Segundo bloque", "Se presentan los trabajos de los grupos 3 y 4. Cada grupo dispondrá de aproximadamente 15 minutos para la presentación y se destinarán 30 minutos para el intercambio.")

doc.add_heading("Ejes para organizar los intercambios", level=1)
doc.add_paragraph("A partir de los argumentos planteados por quienes participan en esta sesión, se proponen los siguientes ejes para organizar los intercambios.")

axes = [
    (
        "Eje 1 Desafíos de la evaluación de impacto",
        "Aunque ha crecido el uso de evaluaciones cuantitativas para medir trayectorias educativas, persisten barreras institucionales relacionadas con la disponibilidad de datos, la comunicación efectiva de los hallazgos y su articulación con indagaciones cualitativas.",
        [
            ("Retroalimentación efectiva de políticas públicas", "¿En qué medida los resultados de las evaluaciones de impacto generan información valiosa y oportuna para ajustar las intervenciones en el aula?"),
            ("Complementariedad metodológica", "¿Cómo dialogan e integran los resultados cuantitativos de impacto con las evaluaciones de procesos e indagaciones cualitativas para comprender la complejidad escolar?"),
            ("Apropiación y transferencia", "¿Cuáles son las estrategias más fructíferas para comunicar y trabajar los resultados junto con referentes y tomadores de decisión de las políticas educativas?"),
        ],
    ),
    (
        "Eje 2 Integración de fuentes y sentido de uso de la información",
        "La producción fragmentada de datos en circuitos, formatos y tiempos diversos, como evaluaciones, registros administrativos y sistemas nominales, dificulta construir lecturas integrales del sistema educativo.",
        [
            ("Superar la lógica del cruce de bases", "Entender la integración no como un procedimiento puramente técnico, sino como la definición de sentidos de uso: para qué se integra, quién consume la información y qué decisiones habilita o limita."),
            ("Construcción de herramientas para la intervención", "Pasar del diagnóstico generalista a herramientas oportunas, situadas y operativas para todos los niveles educativos."),
        ],
    ),
    (
        "Eje 3 Escalas y temporalidades entre la mirada sistémica y la intervención situada",
        "La tensión entre el seguimiento de la trayectoria individual mediante datos nominales y la comprensión de las tendencias generales del sistema mediante datos agregados.",
        [
            ("Riesgo de despolitización e individualización", "Atender casos puntuales no debe invisibilizar los problemas estructurales de aprendizaje, las desigualdades sociales ni las condiciones institucionales."),
            ("Mediaciones articuladoras", "Diseñar dispositivos técnicos, pedagógicos y ético-políticos que permitan atender la trayectoria real del estudiante sin perder de vista las políticas educativas de alcance sistémico."),
        ],
    ),
    (
        "Eje 4 Evaluación e interpretación pedagógica",
        "El riesgo de convertir el resultado de una evaluación en una medida autosuficiente, aislada y descontextualizada de la calidad educativa.",
        [
            ("Lectura situada", "Cruzar los resultados estandarizados o cuantitativos con las condiciones de producción pedagógicas y socioculturales de cada escuela."),
            ("Mediaciones pedagógicas", "Transformar la difusión de reportes en preguntas de trabajo e interpretaciones útiles para que supervisores, directivos y docentes puedan actuar."),
        ],
    ),
    (
        "Eje 5 Dimensión ético política del dato público y nominalizado",
        "La evaluación como acción política que clasifica, visibiliza, genera discursos y produce consecuencias concretas sobre escuelas, docentes y estudiantes.",
        [
            ("Responsabilidad en la circulación", "Establecer protocolos y criterios éticos para el acceso, uso y difusión de información sensible o nominalizada."),
            ("Prevención de la estigmatización", "Evitar que los sistemas de datos nominales o los rankings implícitos refuercen etiquetas o prejuicios sobre determinadas comunidades educativas."),
        ],
    ),
    (
        "Eje 6 Adaptabilidad de herramientas metodológicas frente al cambio de paradigma educativo",
        "La validez de las métricas tradicionales frente a la transformación del régimen académico y de las trayectorias escolares en Secundaria Aprende.",
        [
            ("Pérdida de vigencia de indicadores clásicos", "Revisar categorías como sobreedad, repitencia y promoción, que pueden dejar de reflejar fielmente la trayectoria estudiantil en los nuevos modelos pedagógicos."),
            ("Flexibilidad y dinamismo de la fórmula", "Rediseñar la operacionalización del índice de logros ITEP-ITES para incorporar métricas emergentes sin perder comparabilidad histórica ni consistencia estructural."),
            ("Definición de nuevos indicadores", "Identificar variables que permitan un seguimiento continuo y cualitativo de los aprendizajes, en lugar de mediciones estáticas del avance de año."),
        ],
    ),
]

for heading, problem, points in axes:
    doc.add_heading(heading, level=2)
    paragraph = doc.add_paragraph()
    paragraph.add_run("Problema central: ").bold = True
    paragraph.add_run(problem)
    for lead, body in points:
        add_bullet(doc, lead, body)

doc.save(OUTPUT)
print(OUTPUT)
