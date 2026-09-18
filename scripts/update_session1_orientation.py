from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "dist/documentos/orientaciones/orientaciones-sesion-1.docx"


def add_bullet(doc: Document, lead: str, body: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(f"{lead}: ")
    run.bold = True
    paragraph.add_run(body)


doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Pt(54)
section.left_margin = section.right_margin = Pt(62)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(20)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
title_border = styles["Title"].element.pPr.find(qn("w:pBdr"))
if title_border is not None:
    styles["Title"].element.pPr.remove(title_border)
styles["Heading 1"].font.name = "Aptos Display"
styles["Heading 1"].font.size = Pt(14)
styles["Heading 1"].font.bold = True
styles["Heading 1"].font.color.rgb = RGBColor(0, 0, 0)
styles["Heading 1"].paragraph_format.space_before = Pt(12)
styles["Heading 1"].paragraph_format.space_after = Pt(5)
styles["Heading 2"].font.name = "Aptos Display"
styles["Heading 2"].font.size = Pt(11.5)
styles["Heading 2"].font.bold = True
styles["Heading 2"].font.color.rgb = RGBColor(0, 0, 0)
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(4)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run("GT 1 Usos de información estadística y acceso a datos públicos")
subtitle = doc.add_paragraph()
subtitle.add_run("Sesión 1 Ejes para la discusión").bold = True
doc.add_paragraph("Moderación a cargo de Lilia V. Toranzos")

doc.add_heading("Trabajos presentados", level=1)
works = [
    ("Nancy Montes", "Acerca de los usos de las evaluaciones estandarizadas: discusiones a partir de un estudio regional comparado."),
    ("Daniel Pinkasz", "Evaluaciones estandarizadas en Argentina: usos y debates. Entre la consolidación y la crítica."),
    ("Gabriela Brandan Zehnder, María José Llanos Pozzi y Marisa Álvarez", "Entre los resultados estandarizados y los sistemas nominales: usos, tensiones y mediaciones de la información educativa para el planeamiento y la gestión de políticas."),
    ("Stella Escandell", "Los Reportes Escuela APRENDER. Algunas consideraciones sobre el acceso a este artefacto, los apoyos para su uso y aspectos problemáticos en el levantamiento de información sobre su uso desde el dispositivo APRENDER."),
    ("Juan Suasnábar", "Estandarizadas pero no cerradas: procesamiento de lenguaje natural y respuestas abiertas en evaluaciones educativas a gran escala."),
]
for author, work in works:
    add_bullet(doc, author, work)

doc.add_heading("Dinámica propuesta para la presentación de trabajos", level=1)
add_bullet(doc, "Primer bloque", "Nancy Montes, Daniel Pinkasz, Gabriela Brandan Zehnder, María José Llanos Pozzi y Marisa Álvarez. Se dispondrá de 10 a 12 minutos por exposición y se reservarán 35 minutos para el intercambio con los asistentes.")
add_bullet(doc, "Segundo bloque", "Stella Escandell y Juan Suasnábar. Se dispondrá de 12 a 15 minutos por exposición y se reservarán 35 minutos para el intercambio con los asistentes.")

doc.add_heading("Ejes que organizan el intercambio", level=1)

axes = [
    (
        "Eje 1 Prescripción de la evidencia vs usos reales y tensión entre datos agregados y sistemas de información nominales",
        "La brecha entre el modelo de políticas basadas en evidencia y los usos efectivos, ceremoniales o de planeamiento de las evaluaciones estandarizadas, sumada a la tensión entre dos regímenes de producción de información: los datos estandarizados agregados y los sistemas de información nominales, situados, oportunos y orientados al seguimiento individual de trayectorias.",
        [
            ("Tensión entre datos estandarizados y sistemas nominales", "Analizar cómo articular los diagnósticos de cohorte a gran escala con los datos de matrícula, asistencia y trayectoria real del SIGED. Discutir el riesgo de que la atención exclusiva en el caso individual diluya la mirada sobre los problemas estructurales y sistémicos de aprendizaje."),
            ("Usos promovidos vs usos efectivos y tipología de uso", "Contrastar las tipologías de uso conceptual y simbólico o ceremonial con la necesidad de construir indicadores de gestión para monitorear políticas estratégicas."),
            ("Gobernanza, sostenibilidad e institucionalidad en escenarios de crisis", "Analizar mecanismos institucionales y normativos, como las mesas federales y el CFE, para sostener la producción de información frente a la volatilidad de las gestiones, los equipos técnicos y la incidencia de agencias internacionales y actores no gubernamentales."),
        ],
    ),
    (
        "Eje 2 Artefactos de información herramientas multiescala y mediaciones para la apropiación pedagógica",
        "La necesidad de construir mediaciones institucionales y técnicas que traduzcan la información estadística y evaluativa en preguntas pedagógicas situadas, superando tanto las fallas de acceso a los reportes nacionales como la falta de acompañamiento técnico en las escuelas.",
        [
            ("Dispositivos de devolución orientados por escala y destinatario", "Analizar herramientas de interfaz según el nivel de decisión, las brechas de acceso a este tipo de información y la necesidad de apoyo técnico."),
            ("Traducción del dato en preguntas pedagógicas e institucionales", "Discutir cómo transformar los reportes en secuencias de lectura que interpelen la dinámica escolar, en lugar de entregar meros consolidados de porcentajes de desempeño. Considerar la mediación de los supervisores y el desafío de favorecer el uso pedagógico y formativo de la información."),
            ("Temporalidad y oportunidad de la intervención", "Considerar la ventaja de contar con señales tempranas de riesgo de abandono o inasistencia reiterada para intervenir antes de que las situaciones de fracaso escolar se consoliden."),
        ],
    ),
    (
        "Eje 3 Sentidos político éticos responsabilidad pública y fronteras metodológicas",
        "Los desafíos éticos y políticos vinculados al manejo de información pública, la sobreestimación de los resultados de evaluación, la articulación entre innovaciones tecnológicas y la responsabilidad sobre los datos nominales sensibles.",
        [
            ("Lectura situada del aporte de las evaluaciones estandarizadas", "Prevenir que las pruebas se conviertan en medidas autosuficientes de la calidad educativa y promover su interpretación cruzada con indicadores contextuales y pedagógicos."),
            ("Responsabilidad ético política en el manejo de datos nominales", "Considerar los recaudos éticos, el acceso responsable y las condiciones de reserva necesarias para evitar la estigmatización o clasificación punitiva de estudiantes y escuelas."),
            ("Algoritmos y prevención de sesgos", "Analizar las implicancias éticas del uso de modelos computacionales, la transparencia de la documentación técnica y la prevención de sesgos en registros lingüísticos heterogéneos."),
            ("Posibilidades y límites de las nuevas herramientas", "Discutir de qué manera la innovación tecnológica puede complementar la interpretación pedagógica y contribuir a políticas orientadas al cumplimiento del derecho a la educación."),
        ],
    ),
]

for heading, problem, points in axes:
    doc.add_heading(heading, level=2)
    paragraph = doc.add_paragraph()
    paragraph.add_run("Problema central: ").bold = True
    paragraph.add_run(problem)
    for lead, body in points:
        add_bullet(doc, lead, body)

doc.save(OUTPUT)
print(OUTPUT)
