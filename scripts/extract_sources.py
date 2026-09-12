from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(r"D:\89_SAIE\01_GT")
OUT = Path("tmp/source_extract")


def extract_docx(path: Path) -> dict:
    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    tables = []
    for table in doc.tables:
        tables.append([[cell.text.strip() for cell in row.cells] for row in table.rows])
    return {"paragraphs": paragraphs, "tables": tables}


def extract_xlsx(path: Path) -> dict:
    wb = load_workbook(path, data_only=False, read_only=True)
    result = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = [value for value in row]
            if any(value not in (None, "") for value in values):
                rows.append(values)
        result[ws.title] = rows
    return result


def extract_pdf(path: Path) -> dict:
    reader = PdfReader(path)
    return {"pages": [(page.extract_text() or "").strip() for page in reader.pages]}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    inventory = []
    for path in sorted(ROOT.rglob("*")):
        if not path.is_file():
            continue
        relative = path.relative_to(ROOT)
        inventory.append({"path": str(relative), "bytes": path.stat().st_size})
        suffix = path.suffix.lower()
        if suffix == ".docx":
            payload = extract_docx(path)
        elif suffix == ".xlsx":
            payload = extract_xlsx(path)
        elif suffix == ".pdf":
            payload = extract_pdf(path)
        else:
            continue
        target = OUT / ("__".join(relative.parts) + ".json")
        target.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    (OUT / "inventory.json").write_text(
        json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8"
    )


if __name__ == "__main__":
    main()
